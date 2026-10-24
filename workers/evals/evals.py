#!/usr/bin/env python3
"""Thin CLI over the vendored kit's evals and tools. JSON in, JSON out. Never calls Claude.

    python3 workers/evals/evals.py run <suite> --course <dir> [--ids a b] [--label L] [--json]

Suites
  config    tools/config_check.py                      -> {ok, errors[], warnings[]}
  static    eval/static_eval.py                        -> {lessons:{id:{pass, errors, warnings, ...}}}
  e2e taste overlap                                    -> {lessons:{id:{pass, problems|K*|O*}}}
  suite     static + e2e + taste + overlap, 3 shards   -> {lessons:{id:{static, e2e, taste, overlap, pass}}}
  lx        eval/lx_auto.py [--apply-minutes]          -> {lessons:{id:{grade, est_min, acr, ...}}, means}
  quote     tools/quote_check.py --raw F [--marker M]  -> {ok, missing[]}
  urls      every URL in src NOTES resolves            -> {ok, urls:[{url, status}]}
  capture   eval/capture.py                            -> {shots:{id:[png paths]}}
  diff      tools/diff_added.py --snapshot DIR         -> {lessons:{id:{added:n, lines[]}}}
  build     build.py                                   -> {out, bytes, missing[]}
  patterns  autoplay right+wrong at 1300 and 390 px    -> {patterns:{n:{pass, runs[]}}}
  smoke     the kit's own examples/cissp parity run    -> {ok, steps:{...}}
  convert   --in FILE --out FILE.md                    -> {ok, chars, format}
  classify  --in FILE.md [--domain D]                  -> {kind, confidence, signals, domain, mismatch}
"""
import argparse, glob, json, os, re, shutil, subprocess, sys, tempfile, time
from concurrent.futures import ThreadPoolExecutor

HERE = os.path.dirname(os.path.abspath(__file__))
STUDIO = os.path.dirname(os.path.dirname(HERE))
KIT = os.path.join(STUDIO, 'vendor', 'kit')


def sh(cmd, cwd, env=None, timeout=1800):
    e = dict(os.environ); e.update(env or {})
    t0 = time.time()
    p = subprocess.run(cmd, cwd=cwd, env=e, capture_output=True, text=True, timeout=timeout)
    return p.returncode, (p.stdout or '') + (p.stderr or ''), round(time.time() - t0, 1)


def load(path, default=None):
    try:
        return json.load(open(path))
    except Exception:
        return default if default is not None else {}


def order(course):
    return load(os.path.join(course, 'course.json')).get('order', [])


def all_ids(course, ids):
    if ids: return ids
    have = {os.path.basename(f)[:-3] for f in glob.glob(os.path.join(course, 'lessons', '*.js'))}
    return [i for i in order(course) if i in have] or sorted(have)


def tail(out, n=40):
    return '\n'.join(out.strip().split('\n')[-n:])


# ---------------------------------------------------------------- kit wrappers
def s_config(a):
    code, out, sec = sh(['python3', 'tools/config_check.py'], a.course)
    errs = [l[6:] for l in out.splitlines() if l.startswith('ERROR ')]
    warns = [l[6:] for l in out.splitlines() if l.startswith('warn  ')]
    return {'ok': code == 0, 'errors': errs, 'warnings': warns, 'seconds': sec}


def s_static(a):
    ids = all_ids(a.course, a.ids)
    code, out, sec = sh(['python3', 'eval/static_eval.py'] + ids, a.course)
    rep = load(os.path.join(a.course, 'eval', 'results', 'static.json'))
    return {'ok': code == 0, 'lessons': {i: rep.get(i, {'pass': False, 'errors': ['no result']}) for i in ids},
            'log': tail(out), 'seconds': sec}


def _browser(kind, a, ids, label):
    script = {'e2e': 'eval/e2e.py', 'taste': 'eval/taste_auto.py', 'overlap': 'eval/overlap.py'}[kind]
    env = {'e2e': 'E2E_OUT', 'taste': 'TASTE_OUT', 'overlap': 'OVERLAP_OUT'}[kind]
    out_file = os.path.join(a.course, 'eval', 'results', f'{kind}-{label}.json')
    if os.path.exists(out_file): os.remove(out_file)
    code, out, sec = sh(['python3', script] + ids, a.course, {env: out_file})
    rep = load(out_file)
    res = {}
    for i in ids:
        r = rep.get(i)
        if r is None:
            res[i] = {'pass': False, 'problems': ['harness: no result'], 'log': tail(out, 12)}
        else:
            res[i] = r
    return {'ok': code == 0 and all(r.get('pass') for r in res.values()), 'lessons': res, 'log': tail(out), 'seconds': sec}


def s_e2e(a): return _browser('e2e', a, all_ids(a.course, a.ids), a.label)
def s_taste(a): return _browser('taste', a, all_ids(a.course, a.ids), a.label)
def s_overlap(a): return _browser('overlap', a, all_ids(a.course, a.ids), a.label)


def s_suite(a):
    """The kit's tools/suite.sh, with ids and JSON: static, then 3 shards of e2e -> taste -> overlap."""
    ids = all_ids(a.course, a.ids)
    t0 = time.time()
    st = s_static(argparse.Namespace(**{**vars(a), 'ids': ids}))
    n = max(1, (len(ids) + 2) // 3)
    shards = [ids[k:k + n] for k in range(0, len(ids), n)]

    def shard(k, part):
        lab = f'{a.label}{"abc"[k]}'
        return {kind: _browser(kind, a, part, lab) for kind in ('e2e', 'taste', 'overlap')}

    with ThreadPoolExecutor(max_workers=3) as ex:
        results = list(ex.map(lambda kv: shard(*kv), enumerate(shards)))
    lessons = {}
    for i in ids:
        row = {'static': st['lessons'].get(i, {})}
        for r in results:
            for kind in ('e2e', 'taste', 'overlap'):
                if i in r[kind]['lessons']: row[kind] = r[kind]['lessons'][i]
        row['pass'] = all(row.get(k, {}).get('pass') for k in ('static', 'e2e', 'taste', 'overlap'))
        lessons[i] = row
    return {'ok': all(r['pass'] for r in lessons.values()), 'lessons': lessons, 'seconds': round(time.time() - t0, 1)}


def s_lx(a):
    code, out, sec = sh(['python3', 'eval/lx_auto.py'], a.course)
    rep = load(os.path.join(a.course, 'eval', 'results', 'lx_auto.json'))
    applied = {}
    if a.apply_minutes:
        for i, r in rep.items():
            p = os.path.join(a.course, 'lessons', i + '.js')
            if not os.path.exists(p) or 'est_min' not in r: continue
            m = max(1, round(r['est_min']))
            s = open(p).read()
            s2, k = re.subn(r'(\bminutes\s*:\s*)\d+', lambda mm: mm.group(1) + str(m), s, count=1)
            if k and s2 != s: open(p, 'w').write(s2); applied[i] = m
    means = [l for l in out.splitlines() if l.startswith('means:') or l.startswith('targets:')]
    return {'ok': code == 0, 'lessons': rep, 'summary': means, 'applied_minutes': applied, 'seconds': sec}


def s_quote(a):
    cmd = ['python3', 'tools/quote_check.py', a.raw] + ([a.marker] if a.marker else [])
    code, out, sec = sh(cmd, a.course)
    missing = [l for l in out.splitlines() if 'NOT IN SOURCE' in l]
    return {'ok': code == 0, 'missing': missing, 'log': tail(out, 10), 'seconds': sec}


def s_urls(a):
    import urllib.request, ssl
    urls = {}
    for f in sorted(glob.glob(os.path.join(a.course, 'src', '*.md'))):
        t = open(f).read()
        notes = t.split('NOTES', 1)[1] if 'NOTES' in t else ''
        for u in re.findall(r'https?://[^\s)>\]]+', notes): urls.setdefault(u.rstrip('.,;'), []).append(os.path.basename(f))
    ctx = ssl.create_default_context(cafile=os.environ.get('SSL_CERT_FILE') or None)

    def check(u):
        for method in ('HEAD', 'GET'):
            try:
                req = urllib.request.Request(u, method=method, headers={'User-Agent': 'trust-course-studio/1 (link check)'})
                with urllib.request.urlopen(req, timeout=15, context=ctx) as r: return r.status
            except Exception as e:
                code = getattr(e, 'code', None)
                if code and code < 400: return code
                if method == 'GET': return code or str(e)[:80]
    with ThreadPoolExecutor(max_workers=8) as ex:
        res = dict(zip(urls, ex.map(check, urls)))
    def kind(st):
        if isinstance(st, int): return 'ok' if st < 400 else 'broken'
        return 'unreachable'  # proxy refusal, timeout, DNS: this machine cannot tell, a person must look
    rows = [{'url': u, 'status': res[u], 'files': urls[u], 'ok': kind(res[u]) == 'ok', 'kind': kind(res[u])} for u in urls]
    return {'ok': all(r['ok'] for r in rows), 'urls': rows}


def s_capture(a):
    ids = all_ids(a.course, a.ids)
    code, out, sec = sh(['python3', 'eval/capture.py', '--label', a.label, '--ids'] + ids, a.course, timeout=3600)
    base = os.path.join(a.course, 'eval', 'taste', a.label)
    shots = {i: sorted(os.path.relpath(p, a.course) for p in glob.glob(os.path.join(base, i, '*.png'))) for i in ids}
    return {'ok': code == 0, 'shots': shots, 'log': tail(out, 10), 'seconds': sec}


def s_diff(a):
    code, out, sec = sh(['python3', 'tools/diff_added.py', a.snapshot], a.course)
    d = os.path.join(a.course, 'eval', 'results', 'diffs')
    lessons = {}
    for f in sorted(glob.glob(os.path.join(d, '*.added.txt'))):
        lines = [l for l in open(f).read().splitlines() if l.strip()]
        lessons[os.path.basename(f)[:-10]] = {'added': len(lines), 'lines': lines}
    return {'ok': code == 0, 'lessons': lessons, 'seconds': sec}


def s_build(a):
    code, out, sec = sh(['python3', 'build.py'], a.course)
    m = re.search(r'(\S+index\.html) (\d+) bytes; missing: (.*)', out)
    return {'ok': code == 0, 'out': m.group(1) if m else None, 'bytes': int(m.group(2)) if m else 0,
            'missing': eval(m.group(3)) if m else [], 'log': tail(out, 5)}


# ---------------------------------------------------------------- pattern contract
def s_patterns(a):
    """Each pattern page: autoplay right and wrong paths at 1300 and 390 px. 4 of 4 runs must pass:
    no page errors, the gate message fires, no sideways scroll."""
    from playwright.sync_api import sync_playwright
    pdir = os.path.abspath(a.pages) if a.pages else os.path.join(a.course, 'patterns')
    files = sorted(glob.glob(os.path.join(pdir, '*.html')))
    if a.ids: files = [f for f in files if os.path.basename(f)[:-5] in a.ids]
    out = {}
    with sync_playwright() as p:
        br = p.chromium.launch()
        for f in files:
            n = os.path.basename(f)[:-5]; runs = []
            for width in (1300, 390):
                for path in ('right', 'wrong'):
                    ctx = br.new_context(viewport={'width': width, 'height': 900})
                    ctx.add_init_script('window.__FAST=true; window.__gate=null; addEventListener("message", e => { if (e.data && e.data.type === "trust-pattern") window.__gate = e.data.event; });')
                    pg = ctx.new_page(); errs = []
                    pg.on('pageerror', lambda e: errs.append(str(e)))
                    r = {'width': width, 'path': path}
                    try:
                        pg.goto('file://' + f); pg.wait_for_timeout(100)
                        if not pg.evaluate('typeof window.__autoplay === "function"'):
                            errs.append('no window.__autoplay')
                        else:
                            pg.evaluate(f'window.__autoplay("{path}")')
                            pg.wait_for_function('window.__gate !== null', timeout=8000)
                        r['gate'] = pg.evaluate('window.__gate')
                        r['overflow'] = pg.evaluate('document.documentElement.scrollWidth - innerWidth')
                        if a.shots:
                            sd = os.path.join(a.course, 'eval', 'patterns'); os.makedirs(sd, exist_ok=True)
                            pg.screenshot(path=os.path.join(sd, f'{n}-{width}-{path}.png'), full_page=True)
                    except Exception as e:
                        errs.append(str(e).split('\n')[0][:160])
                    r['errors'] = errs
                    r['pass'] = not errs and bool(r.get('gate')) and r.get('overflow', 0) <= 1
                    runs.append(r); ctx.close()
            out[n] = {'pass': all(r['pass'] for r in runs), 'passed': sum(r['pass'] for r in runs), 'runs': runs}
        br.close()
    return {'ok': all(v['pass'] for v in out.values()), 'patterns': out}


# ---------------------------------------------------------------- parity smoke
def s_smoke(a):
    """Copy examples/cissp into a scratch course with the PATCHED kit and run the kit's own checks."""
    tmp = tempfile.mkdtemp(prefix='tcs-smoke-')
    try:
        for d in ('engine', 'eval', 'tools', 'examples'): shutil.copytree(os.path.join(KIT, d), os.path.join(tmp, d))
        for f in ('build.py', 'AUTHORING.md'): shutil.copy(os.path.join(KIT, f), tmp)
        if not a.pristine:
            code, out, _ = sh(['patch', '-p1', '-s', '-i', os.path.join(STUDIO, 'vendor', 'patches', 'pattern-block.diff')], tmp)
            if code: return {'ok': False, 'steps': {'patch': {'ok': False, 'log': out}}}
        for x in os.listdir(os.path.join(KIT, 'examples', 'cissp')):
            s = os.path.join(KIT, 'examples', 'cissp', x)
            (shutil.copytree if os.path.isdir(s) else shutil.copy)(s, os.path.join(tmp, x))
        ns = argparse.Namespace(course=tmp, ids=None, label='smoke', apply_minutes=False)
        steps = {'config': s_config(ns)}
        steps['build'] = s_build(ns)
        steps['suite'] = s_suite(ns)
        steps['lx'] = s_lx(ns)
        for k in ('suite',):  # keep the report small
            steps[k] = {'ok': steps[k]['ok'], 'lessons': {i: r['pass'] for i, r in steps[k]['lessons'].items()}, 'seconds': steps[k]['seconds']}
        steps['lx'] = {'ok': steps['lx']['ok'], 'summary': steps['lx']['summary']}
        return {'ok': all(s['ok'] for s in steps.values()), 'patched': not a.pristine, 'steps': steps}
    finally:
        if not a.keep: shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- intake
def s_convert(a):
    src, dst = a.inp, a.out
    ext = os.path.splitext(src)[1].lower()
    fmt = ext.lstrip('.')
    if os.path.isdir(src):
        parts = []
        for f in sorted(glob.glob(os.path.join(src, '*'))):
            sub = argparse.Namespace(inp=f, out=dst + '.part')
            r = s_convert(sub)
            if r['ok']: parts.append(f'# {os.path.basename(f)}\n\n' + open(dst + '.part').read())
        if os.path.exists(dst + '.part'): os.remove(dst + '.part')
        open(dst, 'w').write('\n\n'.join(parts)); fmt = 'folder'
    elif ext in ('.md', '.markdown', '.txt'):
        shutil.copy(src, dst)
    elif ext in ('.docx', '.odt', '.html', '.htm', '.rtf', '.epub'):
        code, out, _ = sh(['pandoc', src, '-t', 'gfm', '--wrap=none', '-o', dst], os.path.dirname(dst))
        if code: return {'ok': False, 'error': out[-400:]}
    elif ext == '.pdf':
        code, out, _ = sh(['pdftotext', '-layout', src, dst], os.path.dirname(dst))
        if code: return {'ok': False, 'error': out[-400:]}
    elif ext == '.pptx':
        from pptx import Presentation
        prs = Presentation(src); lines = []
        for k, slide in enumerate(prs.slides, 1):
            lines.append(f'\n## Slide {k}\n')
            for shp in slide.shapes:
                if shp.has_text_frame:
                    for para in shp.text_frame.paragraphs:
                        t = ''.join(r.text for r in para.runs).strip()
                        if t: lines.append(('- ' if para.level else '') + t)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                lines.append('\n> Notes: ' + slide.notes_slide.notes_text_frame.text.strip().replace('\n', ' '))
        open(dst, 'w').write('\n'.join(lines))
    else:
        return {'ok': False, 'error': f'unsupported file type {ext}'}
    text = open(dst).read()
    return {'ok': True, 'format': fmt, 'chars': len(text), 'words': len(text.split()), 'out': dst}


DOMAINS = {
    'ai': ['artificial intelligence', ' ai ', 'model', 'machine learning', 'algorithm', 'chatbot', 'llm', 'training data', 'bias', 'automated decision'],
    'cyber': ['security', 'attack', 'breach', 'malware', 'phishing', 'vulnerability', 'patch', 'encryption', 'access control', 'incident', 'cissp', 'firewall', 'ransomware'],
    'privacy': ['privacy', 'personal data', 'gdpr', 'consent', 'data subject', 'retention', 'dpia', 'controller', 'processor'],
    'data': ['data governance', 'data quality', 'lineage', 'catalogue', 'catalog', 'steward', 'master data', 'metadata'],
    'resilience': ['resilience', 'continuity', 'outage', 'recovery', 'disaster', 'rto', 'rpo', 'failover'],
    'tech': ['change management', 'release', 'deployment', 'architecture', 'technology governance', 'it governance', 'cobit'],
    'third-party': ['third party', 'third-party', 'vendor', 'supplier', 'outsourc', 'due diligence', 'contract'],
}


def s_classify(a):
    t = open(a.inp).read(); low = ' ' + t.lower() + ' '
    words = max(1, len(t.split()))
    sig = {
        'slides': len(re.findall(r'(?m)^##? *Slide \d+', t)),
        'bullets_per_100_lines': round(100 * len(re.findall(r'(?m)^\s*[-*•] ', t)) / max(1, t.count('\n')), 1),
        'dates_per_1k': round(1000 * len(re.findall(r'\b(?:January|February|March|April|May|June|July|August|September|October|November|December|Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\.? \d{1,2}?,? ?\d{4}\b|\b(?:19|20)\d{2}\b', t)) / words, 2),
        'money_per_1k': round(1000 * len(re.findall(r'[$£€]\s?\d[\d,.]*\s?(?:million|billion|m|bn)?', t)) / words, 2),
        'case_words_per_1k': round(1000 * sum(low.count(w) for w in (' court', ' ruled', ' fined', ' settlement', ' testified', ' breach of', ' incident report', ' v. ', ' regulator', ' investigation', ' timeline')) / words, 2),
        'concept_words_per_1k': round(1000 * sum(low.count(w) for w in (' principle', ' concept', ' definition', ' should ', ' must ', ' objective', ' domain', ' control objective', ' best practice', ' framework')) / words, 2),
    }
    case = sig['dates_per_1k'] * 1.2 + sig['money_per_1k'] * 2 + sig['case_words_per_1k'] * 3
    concept = sig['concept_words_per_1k'] * 1.5 + sig['bullets_per_100_lines'] * 0.15 + (4 if sig['slides'] > 5 else 0)
    kind = 'case-rich' if case > concept else 'concept-only'
    conf = round(abs(case - concept) / max(1e-6, case + concept), 2)
    scores = {d: sum(low.count(k) for k in ks) for d, ks in DOMAINS.items()}
    top = max(scores, key=scores.get)
    stated = (a.domain or '').lower()
    stated_key = next((d for d in DOMAINS if d in stated or stated in d), None) if stated else None
    mismatch = bool(stated_key and scores[top] > 0 and stated_key != top and scores.get(stated_key, 0) < 0.35 * scores[top])
    units = re.findall(r'(?m)^#{1,2} +(.+)$', t)
    return {'ok': True, 'kind': kind, 'confidence': conf, 'scores': {'case': round(case, 2), 'concept': round(concept, 2)},
            'signals': sig, 'domain': top, 'domain_scores': scores, 'stated': stated_key, 'mismatch': mismatch,
            'units': units[:80], 'words': words}


SUITES = {k[2:]: v for k, v in globals().items() if k.startswith('s_')}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('cmd', choices=['run']); ap.add_argument('suite', choices=sorted(SUITES))
    ap.add_argument('--course', default='.'); ap.add_argument('--ids', nargs='*'); ap.add_argument('--label', default='studio')
    ap.add_argument('--json', action='store_true'); ap.add_argument('--apply-minutes', action='store_true')
    ap.add_argument('--raw'); ap.add_argument('--marker'); ap.add_argument('--snapshot')
    ap.add_argument('--in', dest='inp'); ap.add_argument('--out'); ap.add_argument('--domain')
    ap.add_argument('--pages'); ap.add_argument('--shots', action='store_true')
    ap.add_argument('--pristine', action='store_true'); ap.add_argument('--keep', action='store_true')
    a = ap.parse_args()
    a.course = os.path.abspath(a.course)
    try:
        res = SUITES[a.suite](a)
    except Exception as e:
        res = {'ok': False, 'error': f'{type(e).__name__}: {e}'}
    res['suite'] = a.suite
    print(json.dumps(res, ensure_ascii=False) if a.json else json.dumps(res, indent=1, ensure_ascii=False))
    sys.exit(0 if res.get('ok') else 1)


if __name__ == '__main__':
    main()
